import io
from datetime import date

import openpyxl
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from .calcs import build_calcs, extract_json
from .checks import run_checks
from .constants import MD_COLUMNS
from .market_data import build_market_data_sheets


_STATUS_FONT = {
    'pass': Font(color='FF42f5b3', bold=True),
    'warn': Font(color='FFf5a742', bold=True),
    'fail': Font(color='FFf54260', bold=True),
}
_HEADER_FILL = PatternFill(fill_type='solid', fgColor='FF1a1a24')
_HEADER_FONT = Font(color='FFe8e8f0', bold=True)


def _autofit(ws, col_widths: list):
    for i, w in enumerate(col_widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w


def build_excel(df_csv, json_data: dict) -> bytes:
    wb = openpyxl.Workbook()

    # ── Sheet 1: Revisión ────────────────────────────────────────────────────
    ws1 = wb.active
    ws1.title = 'Revisión'

    calcs = build_calcs(df_csv)
    jx = extract_json(json_data)
    all_checks = run_checks(calcs, jx)

    headers = ['ODT', 'Grupo', 'Widget / Check', 'Regla', 'Esperado (CSV)', 'JSON actual', 'Estado', 'Δ', 'Página', 'Widget ID']
    ws1.append(headers)
    for cell in ws1[1]:
        cell.fill = _HEADER_FILL
        cell.font = _HEADER_FONT

    for c in all_checks:
        row_data = [c['odt'], c['group'], c['widget'], c['rule'],
                    c['csv'], c['json'], c['status'].upper(), '\n'.join(c['detail']),
                    c.get('page_title', ''), c.get('widget_id', '')]
        ws1.append(row_data)
        rn = ws1.max_row
        ws1.cell(row=rn, column=7).font = _STATUS_FONT.get(c['status'], Font())

    _autofit(ws1, [8, 30, 35, 45, 16, 16, 8, 20, 35, 18])

    # ── Sheet 2: Datos CSV ───────────────────────────────────────────────────
    ws2 = wb.create_sheet('Datos CSV')
    csv_headers = list(df_csv.columns)
    ws2.append(csv_headers)
    for cell in ws2[1]:
        cell.fill = _HEADER_FILL
        cell.font = _HEADER_FONT
    for _, row in df_csv.iterrows():
        ws2.append([str(v) if v is not None else '' for v in row.values])
    _autofit(ws2, [max(12, min(len(h) + 2, 30)) for h in csv_headers])

    # ── Sheets 3-4: Market Data ──────────────────────────────────────────────
    all_records, fondos_records = build_market_data_sheets(df_csv)

    for sheet_name, records in [('look-through', all_records), ('Fondos', fondos_records)]:
        ws = wb.create_sheet(sheet_name)
        ws.append(MD_COLUMNS)
        for cell in ws[1]:
            cell.fill = _HEADER_FILL
            cell.font = _HEADER_FONT
        for rec in records:
            ws.append([rec.get(col, '') if rec.get(col) is not None else '' for col in MD_COLUMNS])
        _autofit(ws, [max(12, min(len(h) + 2, 28)) for h in MD_COLUMNS])

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf.read()
