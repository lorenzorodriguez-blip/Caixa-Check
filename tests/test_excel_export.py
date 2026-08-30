import io

import openpyxl
import pandas as pd

from src.excel_export import build_excel


def _make_inputs():
    df = pd.DataFrame([{'final_market_value': 100, 'custodian': 'Test', 'asset_class_group': 'cash',
                         'sub_asset_class': 'cash', 'isin': 'X', 'asset_description': 'Test asset',
                         'currency': 'eur'}])
    json_data = {'pages': []}
    return df, json_data


def test_build_excel_always_builds_spanish_headers():
    df, json_data = _make_inputs()
    result = build_excel(df, json_data)

    wb = openpyxl.load_workbook(io.BytesIO(result))

    assert wb.sheetnames[0] == 'Revisión'
    ws1 = wb['Revisión']
    header_row = [cell.value for cell in ws1[1]]
    assert header_row == ['ODT', 'Grupo', 'Widget / Check', 'Regla', 'Esperado (CSV)',
                           'JSON actual', 'Estado', 'Δ', 'Página', 'Widget ID']

    assert 'Datos CSV' in wb.sheetnames
    assert 'Fondos' in wb.sheetnames
